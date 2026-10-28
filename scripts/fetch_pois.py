# -*- coding: utf-8 -*-
"""
批量拉取 POI 坐标，生成《点位采集表》草稿（给团队填文案用）

为什么不用"爬"：高德/腾讯都有官方的「地点搜索」接口，返回的坐标就是 GCJ-02，
和小程序地图完全一致，合法、稳定、有免费额度。

用法（先申请一个 key，免费，见 README）：
    # 高德（推荐，个人 5000 次/日）
    python scripts/fetch_pois.py --provider amap --key 你的key --cities 长沙,株洲

    # 腾讯（个人 200 次/日，但一次能拿 20 条，也够）
    python scripts/fetch_pois.py --provider tencent --key 你的key --cities 湘潭

    # 不联网，用内置样例验证流程
    python scripts/fetch_pois.py --selftest

产出：docs/采集草稿-长沙.xlsx 等，结构和《湘潭点位采集表》一致，
      坐标已经填在 F/G 列（状态「未开始」= 坐标是机器给的，等人核对），
      ② 文案补充 留给团队写介绍。
"""
import argparse
import json
import sys
import urllib.parse
import urllib.request
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]

# 城市名 -> 高德 city 参数（adcode）
CITY_CODE = {"湘潭": "430300", "长沙": "430100", "株洲": "430200"}

# 分类 -> 搜索关键词（每个关键词最多翻 2 页，每页 20 条）
KEYWORDS = {
    "景点": ["风景区", "公园", "博物馆", "纪念馆", "寺庙", "古建筑", "美术馆"],
    "红色": ["故居", "革命遗址", "纪念馆"],
    "古镇街区": ["古镇", "老街", "步行街", "历史文化街区"],
    "学府": ["大学", "学院"],
    "美食": ["特色小吃", "老字号", "米粉面馆"],
}

# 明显不该进文旅地图的词
BLOCKLIST = [
    "停车场", "公厕", "卫生间", "派出所", "公交站", "地铁站", "加油站", "充电站",
    "小区", "公寓", "酒店", "宾馆", "民宿", "医院", "卫生院", "银行", "ATM", "超市",
    "便利店", "药店", "幼儿园", "小学", "中学", "职业学院", "职业技术", "培训",
    "公司", "工厂", "仓库", "4S店", "营业厅",
]

RED = "B0392E"
YELLOW = "FFF6DD"
LINE = "D9D0BE"
THIN = Side(style="thin", color=LINE)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
WRAP = Alignment(vertical="center", wrap_text=True)


# ---------------------------------------------------------------- 网络请求
def http_json(url):
    req = urllib.request.Request(url, headers={"User-Agent": "xiangtan-map/1.0"})
    with urllib.request.urlopen(req, timeout=20) as resp:
        return json.loads(resp.read().decode("utf-8"))


def fetch_amap(keyword, city, key, pages=2):
    """高德 地点搜索：返回 [{name, lng, lat, address, type}]"""
    out = []
    for page in range(1, pages + 1):
        params = urllib.parse.urlencode({
            "key": key, "keywords": keyword, "city": CITY_CODE[city],
            "citylimit": "true", "offset": 20, "page": page, "extensions": "base",
        })
        data = http_json("https://restapi.amap.com/v3/place/text?" + params)
        if data.get("status") != "1":
            raise RuntimeError(f"高德接口报错：{data.get('info')}（{data.get('infocode')}）")
        pois = data.get("pois") or []
        for p in pois:
            loc = (p.get("location") or "").split(",")
            if len(loc) != 2:
                continue
            out.append({
                "name": p.get("name", "").strip(),
                "lng": float(loc[0]),
                "lat": float(loc[1]),
                "address": p.get("address") if isinstance(p.get("address"), str) else "",
                "district": p.get("adname") if isinstance(p.get("adname"), str) else "",
                "type": p.get("type", ""),
            })
        if len(pois) < 20:
            break
    return out


def fetch_tencent(keyword, city, key, pages=2):
    """腾讯 地点搜索：返回同样结构"""
    out = []
    for page in range(1, pages + 1):
        params = urllib.parse.urlencode({
            "keyword": keyword, "boundary": f"region({city},0)", "page_size": 20,
            "page_index": page, "key": key,
        })
        data = http_json("https://apis.map.qq.com/ws/place/v1/search?" + params)
        if data.get("status") != 0:
            raise RuntimeError(f"腾讯接口报错：{data.get('message')}（{data.get('status')}）")
        items = data.get("data") or []
        for p in items:
            loc = p.get("location") or {}
            if "lng" not in loc:
                continue
            ad = (p.get("ad_info") or {})
            out.append({
                "name": (p.get("title") or "").strip(),
                "lng": float(loc["lng"]),
                "lat": float(loc["lat"]),
                "address": p.get("address") or "",
                "district": ad.get("district") or "",
                "type": p.get("category") or "",
            })
        if len(items) < 20:
            break
    return out


# ---------------------------------------------------------------- 清洗
def clean(rows):
    """去黑名单、去重（同名近距离算重复）"""
    seen = []
    out = []
    for r in rows:
        name = r["name"]
        if not name or len(name) < 2:
            continue
        if any(bad in name for bad in BLOCKLIST):
            continue
        dup = False
        for s in seen:
            if s["name"] == name and abs(s["lat"] - r["lat"]) < 0.003 and abs(s["lng"] - r["lng"]) < 0.003:
                dup = True
                break
        if dup:
            continue
        seen.append(r)
        out.append(r)
    return out


# ---------------------------------------------------------------- 写表
def write_draft(city, by_category, path):
    """生成的表结构和《湘潭点位采集表》一致：第 3 行表头，第 4 行起是数据，
    这样 import_places.py 可以直接吃。"""
    wb = Workbook()

    def setup_sheet(ws, title, hint, headers):
        n = len(headers)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n)
        c = ws.cell(row=1, column=1, value=title)
        c.font = Font(name="微软雅黑", size=14, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=RED)
        c.alignment = Alignment(vertical="center", indent=1)
        ws.row_dimensions[1].height = 32

        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n)
        c = ws.cell(row=2, column=1, value=hint)
        c.font = Font(name="微软雅黑", size=10, color="736D61")
        c.alignment = Alignment(vertical="center", indent=1)
        ws.row_dimensions[2].height = 22

        for i, (h, w) in enumerate(headers, start=1):
            c = ws.cell(row=3, column=i, value=h)
            c.font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor=RED)
            c.alignment = CENTER
            c.border = BORDER
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.row_dimensions[3].height = 34
        ws.freeze_panes = "D4"

    ws = wb.active
    ws.title = "① 坐标采集"
    ws.sheet_view.showGridLines = False
    setup_sheet(ws, f"① 坐标采集 · {city}草稿（坐标机器抓取，请人工核对）",
                "F/G 是抓到的坐标（高德/腾讯 GCJ-02），核对无误把 I 列改成「已校准」；错的直接改。",
                [("编号", 7), ("点位名称", 22), ("分类", 11), ("旧经度", 13), ("旧纬度", 13),
                 ("✅ 实测经度", 14), ("✅ 实测纬度", 14), ("采集人", 10), ("状态", 10)])

    row, idx = 4, 1
    for cat, items in by_category.items():
        for p in items:
            values = [f"{idx:03d}", p["name"], cat, None, None,
                      round(p["lng"], 6), round(p["lat"], 6), None, "未开始"]
            for col, v in enumerate(values, start=1):
                cell = ws.cell(row=row, column=col, value=v)
                cell.font = Font(name="微软雅黑", size=10)
                cell.border = BORDER
                cell.alignment = WRAP if col == 2 else CENTER
            for col in (6, 7):
                cell = ws.cell(row=row, column=col)
                cell.number_format = "0.000000"
                cell.fill = PatternFill("solid", fgColor=YELLOW)
            row += 1
            idx += 1

    ws2 = wb.create_sheet("② 文案补充")
    ws2.sheet_view.showGridLines = False
    setup_sheet(ws2, f"② 文案补充 · {city}（已有名称和地址，缺介绍）",
                "规则：一句话 ≤30 字；详细介绍 80–120 字；备注或补充自由写。",
                [("编号", 7), ("点位名称", 22), ("一句话介绍（≤30字）", 30),
                 ("详细介绍（80-120字）", 50), ("备注或补充", 34), ("详细地址", 26),
                 ("建议玩多久（分钟）", 12), ("开放时间", 16)])

    row, idx = 4, 1
    for cat, items in by_category.items():
        for p in items:
            addr = ((p.get("district") or "") + (p.get("address") or "")).strip()
            values = [f"{idx:03d}", p["name"], None, None, None, addr, None, None]
            for col, v in enumerate(values, start=1):
                cell = ws2.cell(row=row, column=col, value=v)
                cell.font = Font(name="微软雅黑", size=10)
                cell.border = BORDER
                cell.alignment = WRAP
            row += 1
            idx += 1

    for sheet in (ws, ws2):
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.sheet_properties.pageSetUpPr.fitToPage = True

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# ---------------------------------------------------------------- 样例（不联网）
SAMPLE = {
    "景点": [
        {"name": "橘子洲头", "lng": 112.9581, "lat": 28.1779, "address": "橘子洲景区", "district": "岳麓区", "type": "风景名胜"},
        {"name": "岳麓山", "lng": 112.9385, "lat": 28.1855, "address": "登高路", "district": "岳麓区", "type": "风景名胜"},
    ],
    "红色": [
        {"name": "毛泽东青年艺术雕塑", "lng": 112.9583, "lat": 28.1792, "address": "橘子洲头", "district": "岳麓区", "type": "纪念地"},
    ],
    "古镇街区": [],
    "学府": [
        {"name": "湖南大学", "lng": 112.9347, "lat": 28.1797, "address": "麓山南路", "district": "岳麓区", "type": "学校"},
    ],
    "美食": [],
}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--provider", choices=["amap", "tencent"], default="amap")
    ap.add_argument("--key", default="")
    ap.add_argument("--cities", default="湘潭", help="逗号分隔：湘潭,长沙,株洲")
    ap.add_argument("--out-dir", default=str(ROOT / "docs"))
    ap.add_argument("--pages", type=int, default=2, help="每个关键词翻几页（每页 20 条）")
    ap.add_argument("--selftest", action="store_true", help="不联网，用内置样例跑一遍流程")
    args = ap.parse_args()

    if args.selftest:
        path = write_draft("长沙", SAMPLE, ROOT / "_tmp" / "采集草稿-自测.xlsx")
        total = sum(len(v) for v in SAMPLE.values())
        print(f"✅ 自测通过：生成 {path}（{total} 个点位）")
        print("把 key 给我就能跑真实数据了。")
        return

    if not args.key:
        sys.exit("缺少 --key。去 https://lbs.amap.com/ 免费申请一个「Web服务」key（个人 5000 次/日），见 README。")

    fetch = fetch_amap if args.provider == "amap" else fetch_tencent
    for city in [c.strip() for c in args.cities.split(",") if c.strip()]:
        if city not in CITY_CODE:
            print(f"跳过 {city}：只支持 {list(CITY_CODE)}")
            continue
        by_category, total = {}, 0
        for cat, words in KEYWORDS.items():
            rows = []
            for w in words:
                try:
                    got = fetch(w, city, args.key, args.pages)
                except Exception as e:
                    print(f"  {city}/{w} 失败：{e}")
                    continue
                rows.extend(got)
            by_category[cat] = clean(rows)
            total += len(by_category[cat])
            print(f"  {city} · {cat}: {len(by_category[cat])} 个")
        path = write_draft(city, by_category, Path(args.out_dir) / f"采集草稿-{city}.xlsx")
        print(f"✅ {city}：{total} 个点位 → {path}\n")


if __name__ == "__main__":
    main()
