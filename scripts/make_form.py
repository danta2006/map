# -*- coding: utf-8 -*-
"""
生成 v3 采集表：
  - 坐标不用手工取，已经从高德草稿里填好
  - 团队只需要：① 保留/不保留 ② 补文案 ③ 明显不对的坐标才改
  - 一个工作簿三个 sheet：怎么填 / ① 选点 / ② 文案

用法：python scripts/make_form.py
输出：docs/长株潭点位采集表-v3.xlsx
"""
import os
import json
import shutil
import subprocess
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "长株潭点位采集表-v3.xlsx"
NODE = os.environ.get("NODE_BIN") or shutil.which("node") or "node"

CITY = {"xiangtan": ("湘潭", "采集草稿-湘潭.xlsx"), "changsha": ("长沙", "采集草稿-长沙.xlsx"),
        "zhuzhou": ("株洲", "采集草稿-株洲.xlsx")}
CITY_CN2ID = {v[0]: k for k, v in CITY.items()}
KEY2CN = {"scenic": "景点", "oldtown": "古镇街区", "red": "红色", "school": "学府", "food": "美食"}

# 每个分类保留多少条候选，以及加分关键词
PLAN = {
    "景点": (20, ["国家级", "省级", "风景名胜", "博物馆", "纪念馆", "书院", "寺", "古建", "公园", "湿地", "美术馆"]),
    "红色": (6, ["故居", "纪念馆", "旧址", "遗址"]),
    "古镇街区": (6, ["古镇", "老街", "历史文化街区", "步行街"]),
    "学府": (4, ["大学", "学院"]),
    "美食": (5, ["老字号", "小吃", "米粉"]),
}
DROP = ["停车场", "售票", "入口", "出口", "大门", "游客中心", "服务点", "管理", "办公", "宿舍", "食堂",
        "分店", "便利店", "超市", "小区", "公寓", "酒店", "宾馆", "民宿", "医院", "卫生院", "小学", "中学",
        "幼儿园", "公司", "银行", "药店", "营业厅", "地铁站", "公交站", "充电站"]

RED = "B0392E"
YELLOW = "FFF6DD"
LINE = "D9D0BE"
THIN = Side(style="thin", color=LINE)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
WRAP = Alignment(vertical="center", wrap_text=True)


def load_existing():
    js = "console.log(JSON.stringify(require(String.raw`" + str(ROOT / "miniprogram" / "data" / "places.js") + "`).PLACES))"
    return json.loads(subprocess.run([str(NODE), "-e", js], capture_output=True, text=True, encoding="utf-8").stdout)


def load_draft(city_id):
    p = ROOT / "docs" / CITY[city_id][1]
    if not p.exists():
        return []
    wb = load_workbook(p, data_only=True)
    ws = wb["① 坐标采集"]
    ws2 = wb["② 文案补充"]
    addr = {}
    for r in range(4, ws2.max_row + 1):
        no = ws2.cell(row=r, column=1).value
        a = ws2.cell(row=r, column=6).value
        if no is not None:
            addr[str(no)] = a or ""
    rows = []
    for r in range(4, ws.max_row + 1):
        no = ws.cell(row=r, column=1).value
        name = ws.cell(row=r, column=2).value
        cat = ws.cell(row=r, column=3).value
        lng = ws.cell(row=r, column=6).value
        lat = ws.cell(row=r, column=7).value
        if name and lng and lat:
            rows.append({
                "name": str(name).strip(), "cat": cat, "lng": float(lng), "lat": float(lat),
                "address": addr.get(str(no), "") if no is not None else "",
            })
    return rows


def shortlist(rows, existing_names):
    picked = {}
    for cat, (cap, boosts) in PLAN.items():
        cands = []
        for r in rows:
            if r["cat"] != cat or r["name"] in existing_names:
                continue
            if any(bad in r["name"] for bad in DROP):
                continue
            score = sum(2 for b in boosts if b in r["name"])
            score += 1 if 3 <= len(r["name"]) <= 12 else 0
            cands.append((score, len(r["name"]), r))
        cands.sort(key=lambda x: (-x[0], x[1]))
        picked[cat] = [c[2] for c in cands[:cap]]
    return picked


def main():
    existing = load_existing()
    wb = Workbook()

    # ---------------- 怎么填 ----------------
    ws = wb.active
    ws.title = "怎么填"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 92

    ws.merge_cells("B1:C1")
    c = ws.cell(row=1, column=2, value="长株潭点位采集表 v3 · 坐标已经填好，你只补内容")
    c.font = Font(name="微软雅黑", size=15, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="1D1C19")
    c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 40

    lines = [
        ("坐标不用管", "已用高德「地点搜索」接口批量填好（E/F 列，GCJ-02，和小程序地图同一套坐标）。\n只有当你发现点位明显标错（比如标到江里），才改 E/F 两列，并在 I 列写一句为什么。"),
        ("① 选点", "看「① 选点」这个 sheet：每行一个候选点，「保留」列选 是 / 否 / 待定。\n一个城市留 30–50 个就够，宁缺毋滥：一个点 = 到了能玩、能拍、能吃的地方。"),
        ("② 文案", "看「② 文案」：按编号对上，写一句话（≤30 字）、简介（80–120 字）、备注或补充（自由写，可多行）。\n备注写“本地人才知道”的那句话：几点去、避什么坑、吃什么。"),
        ("不用写的", "分类、城市、地址已经填好；分类要改就在「① 选点」的 D 列下拉改。"),
        ("交回来", "整份表发回，我跑：\n  python scripts/import_places.py --xlsx docs/长株潭点位采集表-v3.xlsx --city changsha\n（--city 换成 xiangtan / zhuzhou / 三个城市各跑一次）"),
    ]
    r = 3
    for k, v in lines:
        a = ws.cell(row=r, column=2, value=k)
        a.font = Font(name="微软雅黑", size=11, bold=True, color=RED)
        a.alignment = Alignment(vertical="top", wrap_text=True)
        b = ws.cell(row=r, column=3, value=v)
        b.font = Font(name="微软雅黑", size=11, color="3F3C35")
        b.alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[r].height = 26 * (1 + v.count("\n"))
        r += 1

    ws.merge_cells(start_row=r + 1, start_column=2, end_row=r + 1, end_column=3)
    tip = ws.cell(row=r + 1, column=2, value="提醒：长沙、株洲现在的点位还很少（14 / 7 个），这次重点是给这两城补量；湘潭已经有 24 个，可以只做微调。")
    tip.font = Font(name="微软雅黑", size=10, color="736D61")
    tip.alignment = Alignment(vertical="center", wrap_text=True, indent=1)

    # ---------------- 数据整理 ----------------
    rows_out = []  # (city_id, name, cat, lng, lat, address, 保留, 备注, summary, intro, note, duration, open)
    for city_id in CITY:
        cn = CITY[city_id][0]
        ex = [p for p in existing if p["city"] == city_id]
        for p in ex:
            rows_out.append({
                "city": city_id, "name": p["name"], "cat": KEY2CN.get(p["category"], p["category"]),
                "lng": p["longitude"], "lat": p["latitude"], "address": p["address"],
                "keep": "是", "why": "", "summary": p["summary"], "intro": p["intro"],
                "note": p.get("note", ""), "duration": p["durationMin"], "open": p["openTime"],
            })
        rows = load_draft(city_id)
        picked = shortlist(rows, {p["name"] for p in ex})
        for cat, items in picked.items():
            for it in items:
                rows_out.append({
                    "city": city_id, "name": it["name"], "cat": cat,
                    "lng": it["lng"], "lat": it["lat"], "address": it.get("address", ""),
                    "keep": "待定", "why": "", "summary": "", "intro": "", "note": "",
                    "duration": "", "open": "",
                })
        print(f"{cn}: 已有 {len(ex)} + 候选 {sum(len(v) for v in picked.values())}")

    # ---------------- ① 选点 ----------------
    ws1 = wb.create_sheet("① 选点")
    heads1 = [("编号", 6), ("城市", 7), ("名称", 26), ("分类", 11), ("经度", 13), ("纬度", 13),
              ("地址", 30), ("保留", 8), ("备注（为什么改/为什么留）", 26)]
    for i, (h, w) in enumerate(heads1, start=1):
        c = ws1.cell(row=1, column=i, value=h)
        c.font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=RED)
        c.alignment = CENTER
        c.border = BORDER
        ws1.column_dimensions[get_column_letter(i)].width = w
    ws1.row_dimensions[1].height = 32

    for idx, p in enumerate(rows_out, start=1):
        row = idx + 1
        vals = [f"{idx:03d}", CITY[p["city"]][0], p["name"], p["cat"],
                round(p["lng"], 6), round(p["lat"], 6), p["address"], p["keep"], p["why"]]
        for col, v in enumerate(vals, start=1):
            c = ws1.cell(row=row, column=col, value=v)
            c.font = Font(name="微软雅黑", size=10)
            c.border = BORDER
            c.alignment = WRAP if col in (3, 7, 9) else CENTER
            if col in (5, 6):
                c.number_format = "0.000000"
        ws1.row_dimensions[row].height = 26

    last = len(rows_out) + 1
    ws1.freeze_panes = "C2"
    ws1.auto_filter.ref = f"A1:I{last}"
    dv_keep = DataValidation(type="list", formula1='"是,否,待定"', allow_blank=False)
    dv_cat = DataValidation(type="list", formula1='"景点,古镇街区,红色,学府,美食"', allow_blank=False)
    ws1.add_data_validation(dv_keep)
    ws1.add_data_validation(dv_cat)
    dv_keep.add(f"H2:H{last}")
    dv_cat.add(f"D2:D{last}")
    for val, color in (("是", "E3F1E8"), ("否", "F0EDE6"), ("待定", "FBE6E2")):
        ws1.conditional_formatting.add(
            f"H2:H{last}", FormulaRule(formula=[f'$H2="{val}"'], fill=PatternFill("solid", fgColor=color)))

    # ---------------- ② 文案 ----------------
    ws2 = wb.create_sheet("② 文案")
    heads2 = [("编号", 6), ("名称", 26), ("一句话（≤30字）", 30), ("简介（80–120字）", 52),
              ("备注或补充", 36), ("建议时长(分钟)", 12), ("开放时间", 16)]
    for i, (h, w) in enumerate(heads2, start=1):
        c = ws2.cell(row=1, column=i, value=h)
        c.font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=RED)
        c.alignment = CENTER
        c.border = BORDER
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.row_dimensions[1].height = 32

    for idx, p in enumerate(rows_out, start=1):
        row = idx + 1
        vals = [f"{idx:03d}", p["name"], p["summary"], p["intro"], p["note"], p["duration"], p["open"]]
        for col, v in enumerate(vals, start=1):
            c = ws2.cell(row=row, column=col, value=v if v != "" else None)
            c.font = Font(name="微软雅黑", size=10)
            c.border = BORDER
            c.alignment = CENTER if col in (1, 6) else WRAP
        ws2.row_dimensions[row].height = 52 if p["intro"] else 30

    ws2.freeze_panes = "C2"
    ws2.auto_filter.ref = f"A1:G{last}"

    for sheet in (ws, ws1, ws2):
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.sheet_properties.pageSetUpPr.fitToPage = True

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"✅ 生成 {OUT}（{len(rows_out)} 行：湘潭 {sum(1 for r in rows_out if r['city']=='xiangtan')}，"
          f"长沙 {sum(1 for r in rows_out if r['city']=='changsha')}，株洲 {sum(1 for r in rows_out if r['city']=='zhuzhou')}）")


if __name__ == "__main__":
    main()
